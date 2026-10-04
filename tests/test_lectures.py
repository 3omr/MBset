"""Tests for `mbset.py lectures` (plan → match → apply → check). Synthetic files in temp dirs only.

Run:  python3 -m unittest discover -s tests -v
"""

from __future__ import annotations

import argparse
import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

SCRIPTS = next(p for p in (Path(__file__).resolve().parents[1] / d / "mbset-module-curator/scripts"
                for d in (".agents/skills", "skills")) if p.exists())
sys.path.insert(0, str(SCRIPTS))

from mbset import lectures  # noqa: E402
from mbset.common import sha256  # noqa: E402

try:
    import fitz  # noqa: F401
    HAVE_FITZ = True
except ImportError:  # pragma: no cover
    HAVE_FITZ = False

PREFIX = "DamiettaFa_TESTMOD_1_"


def cli(*argv: str) -> tuple[int, str]:
    ap = argparse.ArgumentParser(prog="mbset.py")
    sub = ap.add_subparsers(dest="cmd", required=True)
    lectures.register(sub)
    args = ap.parse_args(list(argv))
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = args.fn(args)
    return code, buf.getvalue()


def make_pdf(path: Path, pages: int = 1, text: str = "lecture text " * 30) -> Path:
    import fitz
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    for n in range(pages):
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), f"{text} page {n + 1}")
    doc.save(path)
    doc.close()
    return path


def make_export(path: Path, rows: list[tuple[str, str, str | None]]) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Subcategories"
    ws.append(lectures.HEADERS)
    for n, (sid, name, tag) in enumerate(rows, 1):
        ws.append([PREFIX[:-1], "TestMod", sid, name, tag, None, None, None, None, None, None, 0, n * 10])
    wb.save(path)
    return path


ROWS = [
    (PREFIX + "HYPERTHYROIDISM", "Hyperthyroidism", "Internal Medicine"),
    (PREFIX + "HYPOTHYROIDISM", "Hypothyroidism", "Internal Medicine"),
    (PREFIX + "CONGENITAL_HYPOTHYROIDISM_CRETINISM", "Congenital hypothyroidism (cretinism)", "Pediatric"),
    (PREFIX + "HASHIMOTO_S_THYROIDITIS", "Hashimoto's Thyroiditis", "Internal Medicine"),
    (PREFIX + "CALCIUM_HOMEOSTASIS", "Bio: Calcium homeostasis", "Integrated Sessions"),
    (PREFIX + "DYSLIPIDEMIA", "Dyslipidemia", "Internal Medicine"),
]


class MatchingTests(unittest.TestCase):
    def test_opposite_prefixes_never_match(self):
        self.assertEqual(lectures._tok("hyperthyroidism", "hypothyroidism"), 0.0)
        self.assertEqual(lectures._tok("hyperthyroidism", "hyperparathyroidism"), 0.0)
        self.assertEqual(lectures._tok("microvascular", "macrovascular"), 0.0)
        self.assertGreater(lectures._tok("thyroiditis", "thyrodisis"), 0.75)          # typos still match

    def test_tokens_strip_telegram_prefix_arabic_and_expand_abbreviations(self):
        self.assertEqual(lectures.tokens("0000000332__endocrine htn"), ["endocrine", "hypertension"])
        self.assertEqual(lectures.tokens("0000000376_0377__Posterior_pituitary"), ["posterior", "pituitary"])
        self.assertEqual(lectures.tokens("slides معدل"), ["slides"])

    def test_similarity_orders_candidates(self):
        name = lectures.tokens("Hashimoto's Thyroiditis")
        self.assertGreater(lectures.similarity(name, lectures.tokens("0000000292__Hashimoto thyrodisis")), 0.8)
        self.assertLess(lectures.similarity(name, lectures.tokens("Hypothyroidism")), 0.4)

    def test_classify(self):
        base = Path("/x/Lecture_PowerPoint_PDF_Staging")
        self.assertEqual(lectures.classify(base / "a.pdf", base, None), "powerpoint")
        tg = Path("/x/Telegram_Staging")
        self.assertEqual(lectures.classify(tg / "DocReader" / "a.pdf", tg, None), "docreader")
        self.assertEqual(lectures.classify(tg / "DocReader" / "a.pptx", tg, None), "powerpoint")
        self.assertEqual(lectures.classify(Path("/x/Book/b.pdf"), Path("/x/Book"), None), "book")
        self.assertEqual(lectures.classify(Path("/x/Lectures_OCR_Upload/b.pdf"), Path("/x/Lectures_OCR_Upload"), None),
                         "existing")
        self.assertEqual(lectures.classify(Path("/x/misc/b.pdf"), Path("/x/misc"), "book"), "book")


@unittest.skipUnless(HAVE_FITZ, "PyMuPDF needed")
class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "TestMod"
        self.root.mkdir()
        self.out = Path(self.tmp.name) / "out"

    def tearDown(self):
        self.tmp.cleanup()

    # ---- Phase 1
    def test_plan_from_schedule(self):
        sched = self.root / "schedule.md"
        sched.write_text("# Internal Medicine\n1. Hyperthyroidism\n- Hypothyroidism\n\nPediatric:\n"
                         "* Congenital hypothyroidism (cretinism)\n", encoding="utf-8")
        code, out = cli("lectures", str(self.root), "plan", "--schedule", str(sched))
        self.assertEqual(code, 0)
        self.assertIn("[STOP]", out)
        xlsx = self.root / "subcategories_TestMod.xlsx"
        from openpyxl import load_workbook
        ws = load_workbook(xlsx)["Subcategories"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), lectures.HEADERS)
        self.assertEqual([r[3] for r in rows[1:]], ["Hyperthyroidism", "Hypothyroidism",
                                                    "Congenital hypothyroidism (cretinism)"])
        self.assertEqual([r[4] for r in rows[1:]], ["Internal Medicine", "Internal Medicine", "Pediatric"])
        self.assertTrue(all(r[0] is None and r[2] is None for r in rows[1:]))   # ids empty
        self.assertEqual([r[12] for r in rows[1:]], [10, 20, 30])
        self.assertEqual({r[1] for r in rows[1:]}, {"TestMod"})
        code, out = cli("lectures", str(self.root), "plan", "--schedule", str(sched))
        self.assertEqual(code, 1)                                                # no silent overwrite
        # the Phase-1 file is refused where an export is needed
        with self.assertRaises(SystemExit):
            lectures.read_export(xlsx)

    def test_plan_from_dir_detects_parts(self):
        d = self.root / "raw"
        for rel in ("Anatomy/Joints Part 1.pdf", "Anatomy/Joints Part 2.pdf", "Anatomy/Epithelium.pptx",
                    "Physiology/02_Action potential.pdf"):
            (d / rel).parent.mkdir(parents=True, exist_ok=True)
            (d / rel).write_bytes(b"x")
        code, out = cli("lectures", str(self.root), "plan", "--from-dir", str(d), "--out", str(self.out / "p.xlsx"))
        self.assertEqual(code, 0)
        self.assertIn("multi-part", out)
        from openpyxl import load_workbook
        names = [r[3] for r in list(load_workbook(self.out / "p.xlsx").active.iter_rows(values_only=True))[1:]]
        self.assertEqual(names, ["Epithelium", "Joints Part 1", "Joints Part 2", "Action potential"])
        cli("lectures", str(self.root), "plan", "--from-dir", str(d), "--out", str(self.out / "m.xlsx"), "--merge-parts")
        names = [r[3] for r in list(load_workbook(self.out / "m.xlsx").active.iter_rows(values_only=True))[1:]]
        self.assertEqual(names, ["Epithelium", "Joints", "Action potential"])

    # ---- Phase 2
    def _sources(self):
        r = self.root
        make_pdf(r / "Lecture_PowerPoint_PDF_Staging/0000000315__Hypothyroidism.pdf", 3)
        make_pdf(r / "Lecture_PowerPoint_PDF_Staging/0000000294__Congenital Hypothyroidism.pdf", 2)
        make_pdf(r / "Lecture_PowerPoint_PDF_Staging/0000000292__Hashimoto thyrodisis.pdf", 2)
        (r / "Telegram_Staging/DocReader").mkdir(parents=True)
        (r / "Telegram_Staging/DocReader/0000000292__Hashimoto thyrodisis.pptx").write_bytes(b"pptx twin")
        make_pdf(r / "Telegram_Staging/DocReader/0000000400__hypothyroidism summary.pdf", 1)
        make_pdf(r / "Telegram_Staging/DocReader/0000000401__dyslipidemia.pdf", 1)
        (r / "Telegram_Staging/DocReader/0000000310__Calcium homeostasis.ppt").write_bytes(b"ppt")
        make_pdf(r / "Book/book.pdf", 10)
        make_export(r / "subcategories_TestMod_2026-09-18.xlsx", ROWS)
        (r / "map.json").write_text(json.dumps({
            "Hyperthyroidism": "Book/book.pdf#3-5",                       # book pin, still obeys priority
            PREFIX + "HASHIMOTO_S_THYROIDITIS": "Book/book.pdf#6-6",      # loses to the PowerPoint
            "Bio: Calcium homeostasis": {"files": ["Book/book.pdf#7-8"], "kind": "book", "force": True},
        }), encoding="utf-8")

    def _match(self, *extra):
        return cli("lectures", str(self.root), "match", "--out-dir", str(self.out), "--map", str(self.root / "map.json"),
                   "--sources", str(self.root / "Lecture_PowerPoint_PDF_Staging"), str(self.root / "Telegram_Staging"),
                   "book=" + str(self.root / "Book"), *extra)

    def test_match_priority_pins_and_one_to_one(self):
        self._sources()
        code, out = self._match()
        self.assertEqual(code, 0)
        self.assertFalse((self.root / ".mbset").exists())                    # --out-dir keeps the module clean
        man = json.loads((self.out / "lectures_manifest.json").read_text(encoding="utf-8"))
        by = {r["subcategoryId"][len(PREFIX):]: r for r in man["lectures"]}
        self.assertEqual(by["HYPOTHYROIDISM"]["kind"], "powerpoint")
        self.assertTrue(by["HYPOTHYROIDISM"]["files"][0]["path"].endswith("0000000315__Hypothyroidism.pdf"))
        self.assertTrue(by["CONGENITAL_HYPOTHYROIDISM_CRETINISM"]["files"][0]["path"].endswith("Congenital Hypothyroidism.pdf"))
        self.assertEqual((by["HYPERTHYROIDISM"]["kind"], by["HYPERTHYROIDISM"]["files"][0]["pages"]), ("book", "3-5"))
        hashi = by["HASHIMOTO_S_THYROIDITIS"]
        self.assertEqual(hashi["kind"], "powerpoint")                        # PowerPoint beats the book pin
        self.assertTrue(hashi["files"][0]["path"].endswith(".pdf"))          # the pdf twin beats the pptx
        self.assertFalse(hashi["ambiguous"])
        self.assertEqual((by["CALCIUM_HOMEOSTASIS"]["kind"], by["CALCIUM_HOMEOSTASIS"]["files"][0]["pages"]),
                         ("book", "7-8"))                                    # force beats the .ppt
        self.assertEqual(by["DYSLIPIDEMIA"]["kind"], "docreader")
        self.assertEqual(man["missingCount"], 0)
        self.assertTrue((self.out / "reports/lectures_match.md").exists())
        # the unmatched DocReader hypothyroidism summary is listed for review, the pptx twin is not
        unused = [Path(u["path"]).name for u in man["unusedSources"]]
        self.assertNotIn("0000000292__Hashimoto thyrodisis.pptx", unused)
        code, _ = self._match("--priority", "docreader")                     # configurable priority
        man = json.loads((self.out / "lectures_manifest.json").read_text(encoding="utf-8"))
        by = {r["subcategoryId"][len(PREFIX):]: r for r in man["lectures"]}
        self.assertEqual(by["HYPOTHYROIDISM"]["kind"], "docreader")
        self.assertEqual(by["HASHIMOTO_S_THYROIDITIS"]["status"], "missing")

    def test_apply_and_check(self):
        self._sources()
        (self.root / "map.json").write_text(json.dumps({
            "Hyperthyroidism": "Book/book.pdf#3-5",
            "Bio: Calcium homeostasis": {"files": ["Book/book.pdf#7-8"], "kind": "book", "force": True},
            "Dyslipidemia": None,                                            # forced missing
        }), encoding="utf-8")
        self._match()
        before = {p: sha256(p) for p in self.root.rglob("*") if p.is_file() and "Lectures" != p.parent.name}
        code, out = cli("lectures", str(self.root), "apply", "--out-dir", str(self.out))
        self.assertEqual(code, 0, out)
        lec = self.root / "Lectures"
        self.assertEqual(len(list(lec.glob("*.pdf"))), 5)
        self.assertIn("missing: 60 Dyslipidemia", out)
        import fitz
        with fitz.open(lec / f"{PREFIX}HYPERTHYROIDISM.pdf") as d:
            self.assertEqual(d.page_count, 3)                                # book pages 3-5
        self.assertEqual(sha256(lec / f"{PREFIX}HYPOTHYROIDISM.pdf"),
                         sha256(self.root / "Lecture_PowerPoint_PDF_Staging/0000000315__Hypothyroidism.pdf"))
        for p, h in before.items():                                          # sources untouched
            self.assertTrue(p.exists())
            self.assertEqual(sha256(p), h)
        code, out = cli("lectures", str(self.root), "apply", "--out-dir", str(self.out))
        self.assertIn("written 0", out)                                      # idempotent
        self._match()
        code, out = cli("lectures", str(self.root), "apply", "--out-dir", str(self.out))
        self.assertIn("written 0", out)                                      # … across a re-match too
        # a different file already in place is never overwritten without --force
        make_pdf(lec / f"{PREFIX}HYPOTHYROIDISM.pdf", 1, text="something else")
        code, out = cli("lectures", str(self.root), "apply", "--out-dir", str(self.out))
        self.assertEqual(code, 1)
        self.assertIn("not overwritten", out)
        code, out = cli("lectures", str(self.root), "apply", "--out-dir", str(self.out), "--force")
        self.assertEqual(code, 0)
        # check: one missing row → fail; add it, then a stray → fail; remove stray → pass
        code, out = cli("lectures", str(self.root), "check", "--out-dir", str(self.out))
        self.assertEqual(code, 1)
        self.assertIn("5/6", out)
        make_pdf(lec / f"{PREFIX}DYSLIPIDEMIA.pdf", 1, text="")
        (lec / "old_name.pdf").write_bytes(b"%PDF")
        code, out = cli("lectures", str(self.root), "check", "--out-dir", str(self.out))
        self.assertEqual(code, 1)
        self.assertIn("stray:   old_name.pdf", out)
        (lec / "old_name.pdf").unlink()
        code, out = cli("lectures", str(self.root), "check", "--out-dir", str(self.out))
        self.assertEqual(code, 0, out)
        self.assertIn("text-layer chars", out)                               # the empty PDF is flagged
        self.assertTrue((self.out / "reports/lectures_check.md").exists())

    def test_apply_converts_slides_and_ocr_when_asked(self):
        self._sources()
        (self.root / "map.json").write_text("{}", encoding="utf-8")
        self._match()
        calls = []

        def fake_convert(src, work):
            calls.append(src.name)
            return make_pdf(work / f"{src.stem}.pdf", 2, text="")           # scanned-looking output

        def fake_ocr(src, work):
            calls.append("ocr")
            return make_pdf(work / "ocr_out.pdf", 2)

        with mock.patch.object(lectures, "convert_to_pdf", side_effect=lambda s, w: s if s.suffix == ".pdf"
                               else fake_convert(s, w)), mock.patch.object(lectures, "ocr_pdf", side_effect=fake_ocr):
            code, out = cli("lectures", str(self.root), "apply", "--out-dir", str(self.out), "--ocr")
        self.assertEqual(code, 0, out)
        self.assertIn("0000000310__Calcium homeostasis.ppt", calls)
        self.assertIn("ocr", calls)
        pages, chars = lectures.pdf_stats(self.root / "Lectures" / f"{PREFIX}CALCIUM_HOMEOSTASIS.pdf")
        self.assertEqual(pages, 2)
        self.assertGreater(chars, 0)


if __name__ == "__main__":
    unittest.main()
