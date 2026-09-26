#!/usr/bin/env python3
"""Build one upload-ready lecture PDF per official MBset subcategory.

The mapping is deliberately external and explicit: source pages are never
guessed from filenames or lecture numbers.  Run without --write first; the
script refuses to write unless every official subcategory has a mapping.
"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

import fitz
from openpyxl import load_workbook


MODULE = Path(__file__).resolve().parents[1]
RAW = MODULE / "Lectures_raw" / "Lectures"
DEFAULT_EXPORT = MODULE / "subcategories_NEUROSCIENCE_SYSTEM_2026-09-16.xlsx"
OUT = MODULE / "Lectures"
BUILD_SCRATCH = MODULE.parents[1] / "scratch"


def official_rows(export_path: Path) -> list[dict[str, object]]:
    wb = load_workbook(export_path, read_only=True, data_only=True)
    ws = wb["Subcategories"]
    headers = [cell.value for cell in ws[1]]
    required = {"subcategoryId", "name", "categoryName", "orderIndex"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Official export is missing columns: {sorted(missing)}")
    pos = {name: headers.index(name) for name in required}
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        sid = values[pos["subcategoryId"]]
        if not sid:
            raise ValueError("Official export contains a row without subcategoryId")
        rows.append(
            {
                "subcategoryId": str(sid),
                "name": str(values[pos["name"]] or ""),
                "categoryName": str(values[pos["categoryName"]] or ""),
                "orderIndex": int(values[pos["orderIndex"]] or 0),
            }
        )
    if len({row["subcategoryId"] for row in rows}) != len(rows):
        raise ValueError("Official export contains duplicate subcategoryId values")
    return sorted(rows, key=lambda row: (row["orderIndex"], row["subcategoryId"]))


def read_maps(map_paths: list[Path]) -> list[dict[str, object]]:
    segments: list[dict[str, object]] = []
    for path in map_paths:
        with path.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
        if not isinstance(data, dict) or not isinstance(data.get("segments"), list):
            raise ValueError(f"Invalid lecture map: {path}")
        for index, segment in enumerate(data["segments"], start=1):
            if not isinstance(segment, dict):
                raise ValueError(f"Map item {index} in {path} is not an object")
            for key in ("source", "start", "end", "subcategoryId"):
                if key not in segment:
                    raise ValueError(f"Map item {index} in {path} lacks {key}")
            item = dict(segment)
            item["start"] = int(item["start"])
            item["end"] = int(item["end"])
            item["mapFile"] = str(path)
            segments.append(item)
    return segments


def validate_mapping(
    segments: list[dict[str, object]], rows: list[dict[str, object]]
) -> tuple[dict[str, list[dict[str, object]]], list[str]]:
    official = {str(row["subcategoryId"]): row for row in rows}
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    errors: list[str] = []
    per_source: dict[str, list[tuple[int, int, dict[str, object]]]] = defaultdict(list)

    for segment in segments:
        sid = str(segment["subcategoryId"])
        source = str(segment["source"])
        start, end = int(segment["start"]), int(segment["end"])
        if sid not in official:
            errors.append(f"unknown official subcategoryId: {sid}")
        if start < 1 or end < start:
            errors.append(f"invalid page range {source} p{start}-{end}")
        source_path = RAW / source
        if not source_path.is_file():
            errors.append(f"missing lecture source: {source_path}")
        else:
            page_count = len(fitz.open(source_path))
            if end > page_count:
                errors.append(
                    f"range beyond page count: {source} p{start}-{end} (has {page_count})"
                )
        grouped[sid].append(segment)
        per_source[source].append((start, end, segment))

    for source, ranges in per_source.items():
        ranges.sort()
        for previous, current in zip(ranges, ranges[1:]):
            if current[0] <= previous[1]:
                errors.append(
                    f"overlapping ranges in {source}: p{previous[0]}-{previous[1]} and "
                    f"p{current[0]}-{current[1]}"
                )

    missing = sorted(set(official) - set(grouped), key=lambda sid: official[sid]["orderIndex"])
    if missing:
        errors.append(f"unmapped official subcategories ({len(missing)}): {', '.join(missing)}")
    extra = sorted(set(grouped) - set(official))
    if extra:
        errors.append(f"mapped IDs not in official export: {', '.join(extra)}")
    return grouped, errors


def print_plan(
    grouped: dict[str, list[dict[str, object]]], rows: list[dict[str, object]], errors: list[str]
) -> None:
    by_id = {str(row["subcategoryId"]): row for row in rows}
    print(f"Official subcategories: {len(rows)}")
    print(f"Mapped subcategories:   {len(grouped)}")
    print(f"Mapped page segments:    {sum(len(items) for items in grouped.values())}")
    for row in rows:
        sid = str(row["subcategoryId"])
        items = grouped.get(sid, [])
        ranges = ", ".join(f"{item['source']} p{item['start']}-{item['end']}" for item in items)
        print(f"{row['orderIndex']:>4} | {sid} | {row['name']} | {ranges or 'MISSING'}")
    if errors:
        print("\nPRE-FLIGHT ERRORS:")
        for error in errors:
            print(f"- {error}")


def build(
    grouped: dict[str, list[dict[str, object]]],
    rows: list[dict[str, object]],
    allow_replace: bool,
) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    existing = sorted(path for path in OUT.glob("*.pdf") if path.is_file())
    if existing and not allow_replace:
        raise RuntimeError(
            f"{OUT} already contains {len(existing)} PDFs; pass --allow-replace only after review"
        )

    BUILD_SCRATCH.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix="lecture_build_", dir=str(BUILD_SCRATCH)))
    try:
        for row in rows:
            sid = str(row["subcategoryId"])
            output = fitz.open()
            output.set_metadata(
                {
                    "title": str(row["name"]),
                    "author": "MBset | Assiut | NEUROSCIENCE SYSTEM",
                    "subject": str(row["categoryName"]),
                    "keywords": sid,
                }
            )
            for segment in grouped[sid]:
                source_path = RAW / str(segment["source"])
                with fitz.open(source_path) as source:
                    output.insert_pdf(
                        source,
                        from_page=int(segment["start"]) - 1,
                        to_page=int(segment["end"]) - 1,
                    )
            staged = stage / f"{sid}.pdf"
            output.save(staged, garbage=4, deflate=True)
            output.close()

        if allow_replace:
            for path in existing:
                path.unlink()
        for staged in sorted(stage.glob("*.pdf")):
            staged.replace(OUT / staged.name)
    finally:
        try:
            stage.rmdir()
        except OSError:
            pass


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--map",
        dest="map_paths",
        action="append",
        type=Path,
        required=True,
        help="JSON lecture map; repeat for disjoint map files",
    )
    parser.add_argument("--export", type=Path, default=DEFAULT_EXPORT)
    parser.add_argument("--write", action="store_true", help="write final PDFs after a clean pre-flight")
    parser.add_argument("--allow-replace", action="store_true")
    args = parser.parse_args()

    rows = official_rows(args.export)
    segments = read_maps(args.map_paths)
    grouped, errors = validate_mapping(segments, rows)
    print_plan(grouped, rows, errors)
    if errors:
        return 2
    if not args.write:
        print("\nDry run passed. Re-run with --write to create Lectures/<subcategoryId>.pdf files.")
        return 0
    build(grouped, rows, args.allow_replace)
    print(f"\nWrote {len(rows)} lecture PDFs to {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
