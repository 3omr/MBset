"""Unified reading-order line model for every source format.

`load_lines()` turns a native PDF, an OCRed scan, a .docx, a .pptx or a text file
into a list of `Line` records in true reading order (page 1 col 1, page 1 col 2,
page 2 col 1, ...). Each line keeps the evidence the later stages need:

* ``page``/``bbox`` — where it is, for figure crops and spot-check sheets;
* ``bold``/``color`` — share of bold / coloured characters (marked answers);
* ``marks`` — highlight / ink / box annotations or filled rectangles over it;
* ``conf`` — OCR confidence (``None`` for native text).

Column detection works on any line set (native or OCR): the gutter is the
emptiest vertical band in the middle 40% of the page, excluding header/footer.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from .common import IMAGE_SUFFIXES, Module, clean_inline


@dataclass
class Line:
    text: str
    page: int = 0
    col: int = 0
    bbox: tuple[float, float, float, float] = (0.0, 0.0, 0.0, 0.0)
    bold: float = 0.0
    color: float = 0.0
    marks: list[str] = field(default_factory=list)
    conf: float | None = None
    size: float = 0.0
    split_option: bool = False   # produced by splitting an inline option run

    def to_json(self) -> dict[str, Any]:
        data = asdict(self)
        data["bbox"] = [round(v, 1) for v in self.bbox]
        return data


# --------------------------------------------------------------------------- columns
def detect_split(lines: list[Line], width: float, height: float) -> float | None:
    """Return the gutter x (points) of a two-column page, or None for one column."""
    body = [ln for ln in lines if height * 0.08 < (ln.bbox[1] + ln.bbox[3]) / 2 < height * 0.94
            and ln.text.strip()]
    if len(body) < 8 or width <= 0:
        return None
    bins = 60
    density = [0.0] * bins
    for ln in body:
        x0, x1 = max(ln.bbox[0], 0), min(ln.bbox[2], width)
        a, b = int(x0 / width * bins), int(x1 / width * bins)
        for i in range(max(a, 0), min(b + 1, bins)):
            density[i] += 1
    lo, hi = int(bins * 0.30), int(bins * 0.70)
    gutter = min(range(lo, hi), key=lambda i: (density[i], abs(i - bins // 2)))
    positive = sorted(d for d in density if d > 0)
    median = positive[len(positive) // 2] if positive else 0
    if not median or density[gutter] > 0.25 * median:
        return None
    split = (gutter + 0.5) / bins * width
    left = sum(1 for ln in body if ln.bbox[2] <= split + 2)
    right = sum(1 for ln in body if ln.bbox[0] >= split - 2)
    if left < 0.25 * len(body) or right < 0.25 * len(body):
        return None
    return split


def order_page(lines: list[Line], width: float, height: float, columns: Any = "auto",
               split: float | None = None) -> tuple[list[Line], float | None]:
    """Merge fragments into visual lines and order them column by column."""
    if columns in (1, "1"):
        split = None
    elif split is None:
        split = detect_split(lines, width, height)
        if columns in (2, "2") and split is None:
            split = width / 2
    lines = attach_glyph_marks(drop_margin_numbers(lines, width, height))
    spans: list[Line] = []
    cols: dict[int, list[Line]] = {0: [], 1: []}
    for ln in lines:
        if split is None:
            ln.col = 0
            cols[0].append(ln)
        elif ln.bbox[2] <= split + 4:
            ln.col = 0
            cols[0].append(ln)
        elif ln.bbox[0] >= split - 4:
            ln.col = 1
            cols[1].append(ln)
        else:
            ln.col = 0
            spans.append(ln)
    ordered: list[Line] = []
    top_of_columns = min((ln.bbox[1] for c in cols.values() for ln in c), default=height)
    head = [ln for ln in spans if ln.bbox[3] <= top_of_columns + 2]
    rest = [ln for ln in spans if ln not in head]
    cols[0].extend(rest)
    for group in (head, cols[0], cols[1]):
        ordered.extend(_merge_visual_lines(group))
    return ordered, split


def _merge_visual_lines(group: list[Line]) -> list[Line]:
    group = sorted(group, key=lambda ln: ((ln.bbox[1] + ln.bbox[3]) / 2, ln.bbox[0]))
    merged: list[list[Line]] = []
    for ln in group:
        if merged:
            last = merged[-1][-1]
            h = min(ln.bbox[3] - ln.bbox[1], last.bbox[3] - last.bbox[1]) or 1
            overlap = min(ln.bbox[3], last.bbox[3]) - max(ln.bbox[1], last.bbox[1])
            if overlap > 0.5 * h:
                merged[-1].append(ln)
                continue
        merged.append([ln])
    out = []
    for group_parts in merged:
        group_parts.sort(key=lambda ln: ln.bbox[0])
        # split a visual line at wide horizontal gaps (side-by-side options, margin page numbers),
        # but keep table cells (short tokens like "12", "C") together
        segments: list[list[Line]] = [[group_parts[0]]]
        for ln in group_parts[1:]:
            prev = segments[-1][-1]
            h = max(prev.bbox[3] - prev.bbox[1], 6)
            gap = ln.bbox[0] - prev.bbox[2]
            short = len(prev.text.strip()) <= 4 and len(ln.text.strip()) <= 4
            if gap > 2.5 * h and not short:
                segments.append([ln])
            else:
                segments[-1].append(ln)
        for parts in segments:
            out.append(_join(parts))
    return out


def _join(parts: list[Line]) -> Line:
    if len(parts) == 1:
        return parts[0]
    n = sum(max(len(p.text), 1) for p in parts)
    confs = [p.conf for p in parts if p.conf is not None]
    return Line(
        text=clean_inline(" ".join(p.text for p in parts)),
        page=parts[0].page, col=parts[0].col,
        bbox=(min(p.bbox[0] for p in parts), min(p.bbox[1] for p in parts),
              max(p.bbox[2] for p in parts), max(p.bbox[3] for p in parts)),
        bold=sum(p.bold * max(len(p.text), 1) for p in parts) / n,
        color=sum(p.color * max(len(p.text), 1) for p in parts) / n,
        marks=sorted({m for p in parts for m in p.marks}),
        conf=min(confs) if confs else None,
        size=max(p.size for p in parts),
    )


TICKS = set("✓✔☑🗸\ue73e\ue8fb\ue10b\uf00c\u2705")
CROSSES = set("✗✘☒✕✖\ue711\ue10a\uf00d\u274c")


def attach_glyph_marks(lines: list[Line]) -> list[Line]:
    """Turn tick/cross glyph fragments (Google Forms review icons, ✓/✗) into marks on the row's text."""
    glyphs, rest = [], []
    for ln in lines:
        t = ln.text.strip()
        if t and all(ch in TICKS or ch in CROSSES or ch.isspace() for ch in t):
            glyphs.append(ln)
        else:
            rest.append(ln)
    for g in glyphs:
        kind = "tick" if any(ch in TICKS for ch in g.text) else "cross"
        cy = (g.bbox[1] + g.bbox[3]) / 2
        row = [o for o in rest if o.bbox[1] - 3 <= cy <= o.bbox[3] + 3 and o.page == g.page]
        if row:
            target = min(row, key=lambda o: min(abs(o.bbox[0] - g.bbox[2]), abs(g.bbox[0] - o.bbox[2])))
            target.marks = sorted(set(target.marks) | {kind})
    return rest


def drop_margin_numbers(lines: list[Line], width: float, height: float) -> list[Line]:
    """Remove bare page numbers in a page margin (they would merge into question lines).

    A margin number goes only when nothing sits near it on the same row, so key-grid
    numbers at the page edges survive.
    """
    out = []
    for ln in lines:
        t = ln.text.strip()
        if re.fullmatch(r"(?:page\s*)?\d{1,3}", t, re.I):
            if (ln.bbox[1] > height * 0.92 or ln.bbox[3] < height * 0.07
                    or ln.bbox[0] > width * 0.85 or ln.bbox[2] < width * 0.12):
                cy = (ln.bbox[1] + ln.bbox[3]) / 2
                row = [o for o in lines if o is not ln and o.bbox[1] <= cy <= o.bbox[3]]
                near = min((min(abs(o.bbox[2] - ln.bbox[0]), abs(ln.bbox[2] - o.bbox[0])) for o in row),
                           default=width)
                if near > 0.2 * width:
                    continue
        out.append(ln)
    return out


# --------------------------------------------------------------------------- PDF marks
_ANNOT_KIND = {8: "highlight", 9: "underline", 4: "box", 5: "circle", 15: "ink", 10: "squiggly", 11: "strike"}


def page_mark_rects(page) -> list[tuple[Any, str]]:
    """Rectangles that emphasise text: highlight/ink/box annotations and colour fills."""
    rects = []
    for annot in page.annots() or []:
        kind = _ANNOT_KIND.get(annot.type[0])
        if kind:
            rects.append((annot.rect, kind))
    try:
        drawings = page.get_drawings()
    except Exception:
        drawings = []
    for d in drawings:
        fill = d.get("fill")
        if not fill or d.get("rect") is None:
            continue
        r, g, b = fill[:3]
        if min(r, g, b) > 0.93 or max(r, g, b) < 0.15 or (max(r, g, b) - min(r, g, b)) < 0.12:
            continue  # white, black or grey
        rect = d["rect"]
        if rect.height < 45 and rect.width < page.rect.width * 0.95:
            rects.append((rect, "fill"))
    return rects


def _marks_for(bbox, rects) -> list[str]:
    import fitz

    box = fitz.Rect(bbox)
    area = max(box.get_area(), 1)
    kinds = set()
    for rect, kind in rects:
        inter = fitz.Rect(box) & rect
        if inter.is_empty:
            continue
        if kind in ("ink", "circle", "box"):
            # a tick or circle drawn next to / over the option counts
            if inter.get_area() > 0.02 * area or rect.get_area() < 4 * area:
                kinds.add(kind)
        elif inter.get_area() > 0.35 * area:
            kinds.add(kind)
    return sorted(kinds)


def _near_marks(bbox, rects) -> list[str]:
    """Ink ticks are usually drawn left of the option letter: widen the box leftwards."""
    x0, y0, x1, y1 = bbox
    return _marks_for((x0 - 40, y0 - 2, x1, y1 + 2), [(r, k) for r, k in rects if k in ("ink", "circle", "box")])


def _is_colored(rgb: int) -> bool:
    r, g, b = (rgb >> 16) & 255, (rgb >> 8) & 255, rgb & 255
    return max(r, g, b) - min(r, g, b) > 60 and max(r, g, b) > 90


def native_pdf_pages(path: Path, pages: list[int] | None = None):
    """Yield (page_index, width, height, [Line]) from the PDF text layer."""
    import fitz

    with fitz.open(path) as doc:
        for pno in pages if pages is not None else range(len(doc)):
            page = doc[pno]
            rects = page_mark_rects(page)
            lines = []
            for block in page.get_text("dict")["blocks"]:
                for raw in block.get("lines", []):
                    spans = [s for s in raw["spans"] if s["text"].strip()]
                    if not spans:
                        continue
                    text = clean_inline("".join(s["text"] for s in raw["spans"]))
                    if not text:
                        continue
                    n = sum(len(s["text"].strip()) for s in spans) or 1
                    bold = sum(len(s["text"].strip()) for s in spans
                               if s["flags"] & 16 or re.search(r"bold|black|heavy", s["font"], re.I)) / n
                    color = sum(len(s["text"].strip()) for s in spans if _is_colored(s["color"])) / n
                    bbox = tuple(raw["bbox"])
                    lines.append(Line(text=text, page=pno, bbox=bbox, bold=round(bold, 2),
                                      color=round(color, 2),
                                      marks=sorted(set(_marks_for(bbox, rects)) | set(_near_marks(bbox, rects))),
                                      size=round(max(s["size"] for s in spans), 1)))
            yield pno, page.rect.width, page.rect.height, lines


def split_margin_note(ln: dict[str, Any], gap: float = 24.0, max_conf: float = 0.6) -> dict[str, Any]:
    """Drop handwriting beside a printed line: '… Glyburide   ws gulfa WS'.

    The tail after the widest word gap (> `gap` pt) goes when it is short and poorly recognised —
    printed text is read at > 0.8, handwriting rarely above 0.5.
    """
    words = ln.get("words")
    if not words or len(words) < 2:
        return ln
    gaps = [(words[k + 1][1] - words[k][2], k) for k in range(len(words) - 1)]
    widest, k = max(gaps)
    tail = words[k + 1:]
    if widest < gap or len(tail) > 4 or sum(w[3] for w in tail) / len(tail) > max_conf:
        return ln
    head = words[:k + 1]
    x0, y0, _, y1 = ln["bbox"]
    confs = [w[3] for w in head if w[3] >= 0]
    return dict(ln, text=" ".join(w[0] for w in head), bbox=[x0, y0, head[-1][2], y1],
                conf=round(sum(confs) / len(confs), 3) if confs else ln["conf"],
                margin_note=" ".join(w[0] for w in tail))


def ocr_pages(module: Module, src: dict[str, Any], pages: list[int] | None = None):
    """Yield OCR lines; for PDFs, annotation marks come from the original page."""
    from .common import load_json

    data = load_json(module.ocr_path(src))
    if not data:
        raise SystemExit(f"[-] {src['nn']}: no OCR yet — run `mbset.py ocr <Module> --only {src['nn']}`")
    doc = None
    path = module.source_path(src)
    if path.suffix.lower() == ".pdf":
        import fitz
        doc = fitz.open(path)
    try:
        for pno, pg in enumerate(data["pages"]):
            if pages is not None and pno not in pages:
                continue
            rects = page_mark_rects(doc[pno]) if doc is not None else []
            lines = []
            for ln in pg["lines"]:
                ln = split_margin_note(ln)
                text = clean_inline(ln["text"])
                if not text:
                    continue
                bbox = tuple(ln["bbox"])
                marks = sorted(set(_marks_for(bbox, rects)) | set(_near_marks(bbox, rects))) if rects else []
                lines.append(Line(text=text, page=pno, bbox=bbox, conf=ln["conf"], marks=marks))
            yield pno, pg["width"], pg["height"], lines
    finally:
        if doc is not None:
            doc.close()


def colour_marks(path: Path, lines: list[Line], dpi: int = 72) -> None:
    """Scanned pages: flag lines sitting on a coloured highlight (pixel saturation)."""
    import fitz
    import numpy as np

    by_page: dict[int, list[Line]] = {}
    for ln in lines:
        by_page.setdefault(ln.page, []).append(ln)
    with fitz.open(path) as doc:
        for pno, group in by_page.items():
            pix = doc[pno].get_pixmap(dpi=dpi, annots=True)
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)[..., :3]
            arr = arr.astype(np.int16)
            sat = arr.max(axis=2) - arr.min(axis=2)
            bright = arr.max(axis=2)
            colored = (sat > 70) & (bright > 110)
            scale = dpi / 72.0
            for ln in group:
                x0, y0, x1, y1 = (int(v * scale) for v in ln.bbox)
                patch = colored[max(y0, 0):max(y1, y0 + 1), max(x0, 0):max(x1, x0 + 1)]
                if patch.size and patch.mean() > 0.18:
                    ln.color = max(ln.color, round(float(patch.mean()), 2))
                    if "tint" not in ln.marks:
                        ln.marks = sorted(set(ln.marks) | {"tint"})


# --------------------------------------------------------------------------- docx / pptx
def _docx_numbering(doc) -> dict[tuple[str, str], str]:
    fmt: dict[tuple[str, str], str] = {}
    try:
        numbering = doc.part.numbering_part.element
    except Exception:
        return fmt
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    w = "{%s}" % ns["w"]
    abstract = {}
    for an in numbering.findall("w:abstractNum", ns):
        levels = {}
        for lvl in an.findall("w:lvl", ns):
            nf = lvl.find("w:numFmt", ns)
            levels[lvl.get(w + "ilvl")] = nf.get(w + "val") if nf is not None else "decimal"
        abstract[an.get(w + "abstractNumId")] = levels
    for num in numbering.findall("w:num", ns):
        ref = num.find("w:abstractNumId", ns)
        if ref is None:
            continue
        for ilvl, kind in abstract.get(ref.get(w + "val"), {}).items():
            fmt[(num.get(w + "numId"), ilvl)] = kind
    return fmt


def _fmt_number(kind: str, n: int) -> str:
    if kind in ("lowerLetter", "upperLetter"):
        letter = "abcdefghijklmnopqrstuvwxyz"[(n - 1) % 26]
        return (letter.upper() if kind == "upperLetter" else letter) + ")"
    if kind == "bullet":
        return ""
    return f"{n}."


def docx_lines(path: Path) -> list[Line]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(str(path))
    numbering = _docx_numbering(doc)
    counters: dict[tuple[str, str], int] = {}
    out: list[Line] = []
    y = 0.0

    def emit(par: Paragraph) -> None:
        nonlocal y
        text = clean_inline(par.text)
        if not text:
            return
        prefix = ""
        ppr = par._p.pPr
        if ppr is not None and ppr.numPr is not None and ppr.numPr.numId is not None:
            num_id = str(ppr.numPr.numId.val)
            ilvl = str(ppr.numPr.ilvl.val) if ppr.numPr.ilvl is not None else "0"
            key = (num_id, ilvl)
            counters[key] = counters.get(key, 0) + 1
            for other in list(counters):  # a deeper list restarts under a new parent item
                if other[0] == num_id and int(other[1]) > int(ilvl):
                    counters.pop(other)
            prefix = _fmt_number(numbering.get(key, "decimal"), counters[key])
        runs = [r for r in par.runs if r.text.strip()]
        n = sum(len(r.text.strip()) for r in runs) or 1
        bold = sum(len(r.text.strip()) for r in runs if r.bold or (r.style is not None and r.style.font.bold)) / n
        color = 0.0
        marks = set()
        for r in runs:
            rgb = r.font.color.rgb if r.font.color is not None and r.font.color.type is not None else None
            if rgb is not None and _is_colored(int(str(rgb), 16)):
                color += len(r.text.strip()) / n
            if r.font.highlight_color is not None:
                marks.add("highlight")
            if r.font.underline:
                marks.add("underline")
        y += 12
        out.append(Line(text=f"{prefix} {text}".strip(), page=0, bbox=(0, y, 500, y + 10),
                        bold=round(bold, 2), color=round(color, 2), marks=sorted(marks)))

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            emit(Paragraph(child, doc))
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                for cell in row.cells:
                    for par in cell.paragraphs:
                        emit(par)
    return out


def pptx_lines(path: Path) -> list[Line]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[Line] = []
    for sno, slide in enumerate(prs.slides):
        shapes = sorted((s for s in slide.shapes if s.has_text_frame),
                        key=lambda s: ((s.top or 0), (s.left or 0)))
        for shape in shapes:
            top = (shape.top or 0) / 12700
            for i, par in enumerate(shape.text_frame.paragraphs):
                text = clean_inline("".join(r.text for r in par.runs))
                if not text:
                    continue
                runs = [r for r in par.runs if r.text.strip()]
                n = sum(len(r.text.strip()) for r in runs) or 1
                bold = sum(len(r.text.strip()) for r in runs if r.font.bold) / n
                color = 0.0
                for r in runs:
                    try:
                        rgb = r.font.color.rgb
                    except Exception:
                        rgb = None
                    if rgb is not None and _is_colored(int(str(rgb), 16)):
                        color += len(r.text.strip()) / n
                y = top + i * 14
                out.append(Line(text=text, page=sno, bbox=((shape.left or 0) / 12700, y, 700, y + 12),
                                bold=round(bold, 2), color=round(color, 2)))
    return out


def text_lines(path: Path) -> list[Line]:
    out = []
    for i, raw in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines()):
        text = clean_inline(raw)
        if text:
            out.append(Line(text=text, page=0, bbox=(0, i * 12.0, 500, i * 12.0 + 10)))
    return out


# --------------------------------------------------------------------------- entry point
def load_lines(module: Module, src: dict[str, Any], profile: dict[str, Any]) -> tuple[list[Line], dict[str, Any]]:
    """Return reading-order lines plus layout facts (per-page split, method)."""
    from .common import parse_pages

    path = module.source_path(src)
    suffix = path.suffix.lower()
    mode = profile.get("text") or "auto"
    info: dict[str, Any] = {"method": mode, "splits": {}}
    if mode == "auto":
        if suffix == ".docx":
            mode = "docx"
        elif suffix == ".pptx":
            mode = "pptx"
        elif suffix in (".txt", ".md"):
            mode = "plain"
        elif suffix in IMAGE_SUFFIXES or not src.get("triage", {}).get("has_text"):
            mode = "ocr"
        else:
            mode = "native"
        info["method"] = mode
    if mode == "docx":
        return docx_lines(path), info
    if mode == "pptx":
        return pptx_lines(path), info
    if mode == "plain":
        return text_lines(path), info

    total = src.get("pages") or 1
    pages = parse_pages(profile.get("pages"), total)
    columns = profile.get("columns", "auto")
    forced_split = profile.get("split")
    iterator = native_pdf_pages(path, pages) if mode == "native" else ocr_pages(module, src, pages)
    out: list[Line] = []
    for pno, width, height, lines in iterator:
        ordered, split = order_page(lines, width, height, columns, forced_split)
        info["splits"][pno + 1] = round(split, 1) if split else None
        out.extend(ordered)
    if mode == "ocr" and suffix == ".pdf" and profile.get("answers", {}).get("tint", True):
        colour_marks(path, out)
    return out, info
