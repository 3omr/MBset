"""`mbset.py lectures` — the two-phase lecture/subcategory workflow (AGENTS.md §2.D).

    lectures <Module> plan  --schedule FILE | --from-dir DIR   Phase 1: subcategories_<Module>.xlsx, ids empty → STOP
    lectures <Module> match --sources DIR [KIND=DIR ...]       Phase 2 dry run: official row → lecture source
    lectures <Module> apply [--ocr] [--force]                  copy/convert matches to Lectures/<subcategoryId>.pdf
    lectures <Module> check                                    every row has its PDF, no strays, pages, text layer

`match` writes `.mbset/lectures_manifest.json` + `.mbset/reports/lectures_match.md` (or `--out-dir`).
Sources are never moved or deleted. Source kinds rank by `--priority` (default PowerPoint/College →
Book → DocReader → existing lecture sets); a kind comes from `KIND=DIR`, else from the file extension
(.pptx/.ppt/.odp = powerpoint), else from the folder name. Book page ranges, merges and hand fixes go in
`--map FILE.json`: {"<subcategoryId or name>": "Book/x.pdf#31-32" | ["a.pdf", "b.pptx"] |
{"files": [...], "kind": "book", "force": true} | null (= force missing)}. A pin is a candidate of its
kind and still obeys the priority unless `"force": true`.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from .common import ARABIC, sha256

HEADERS = ["categoryId", "categoryName", "subcategoryId", "name", "tag", "type", "pdf",
           "studyRecommendations", "pdfVersion", "pdfDate", "pdfNote", "questionsCount", "orderIndex"]
SHEET = "Subcategories"
LECTURE_SUFFIXES = {".pdf", ".pptx", ".ppt", ".odp", ".docx", ".doc", ".odt"}
SLIDE_SUFFIXES = {".pptx", ".ppt", ".odp"}
KINDS = ["powerpoint", "book", "docreader", "existing", "other"]
DEFAULT_PRIORITY = KINDS
DIR_KINDS = [  # first hit on any folder name between the source root and the file wins
    ("powerpoint", re.compile(r"power ?point|\bppts?\b|slides?|college", re.I)),
    ("book", re.compile(r"\bbooks?\b", re.I)),
    ("docreader", re.compile(r"docreader|telegram|summar", re.I)),
    ("existing", re.compile(r"lecture|upload|\bocr\b", re.I)),
]
STOP = {"a", "an", "and", "the", "of", "to", "in", "on", "for", "with", "s", "pdf", "copy", "final", "new",
        "lecture", "lec", "dr", "merged"}
# medical prefixes that change the meaning: "hypo…" never fuzzy-matches "hyper…", nor "thyroid…" "parathyroid…"
PREFIXES = ("hyper", "hypo", "para", "pan", "poly", "oligo", "micro", "macro", "pseudo", "peri", "pre", "post",
            "anti", "non", "dys", "intra", "inter", "endo", "exo", "brady", "tachy", "anterior", "posterior")
ABBREV = {"dm": "diabetes mellitus", "htn": "hypertension", "ds": "diseases", "dz": "disease", "ca": "cancer",
          "tb": "tuberculosis", "gn": "glomerulonephritis", "ckd": "chronic kidney disease", "ihd": "ischemic heart disease",
          "mi": "myocardial infarction", "cvs": "cardiovascular", "cns": "nervous", "git": "gastrointestinal",
          "gi": "gastrointestinal", "uti": "urinary infection", "ra": "rheumatoid arthritis", "sle": "lupus",
          "dka": "diabetic ketoacidosis", "t1dm": "type 1 diabetes mellitus", "t2dm": "type 2 diabetes mellitus"}
TG_PREFIX = re.compile(r"^\d{4,}(?:_\d+)*__")
MANIFEST = "lectures_manifest.json"


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


# --------------------------------------------------------------------------- paths
def module_root(arg: str) -> Path:
    root = Path(arg).resolve()
    if not root.is_dir():
        raise SystemExit(f"[-] module folder not found: {root}")
    return root


def meta_dir(root: Path, out_dir: str | None) -> Path:
    """Where manifest + reports live: `--out-dir` or `<Module>/.mbset` (created only on write)."""
    return Path(out_dir).resolve() if out_dir else root / ".mbset"


def _rel(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root))
    except ValueError:
        return str(path.resolve())


def _abs(path: str, root: Path) -> Path:
    """A CLI path: absolute, else relative to the working directory when it exists there, else to the module."""
    p = Path(path)
    if p.is_absolute():
        return p
    return p.resolve() if p.exists() else root / p


# --------------------------------------------------------------------------- export
def find_export(root: Path) -> Path:
    exports = sorted(p for p in root.glob(f"subcategories_{root.name}_*.xlsx") if not p.name.startswith("~$"))
    if not exports:
        raise SystemExit(f"[-] no platform export subcategories_{root.name}_<Date>.xlsx in {root} — "
                         "pass --export (Phase 1 not uploaded yet?)")
    return exports[-1]


def read_export(path: Path, need_ids: bool = True) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it)]
        rows = [dict(zip(headers, r)) for r in it if any(v not in (None, "") for v in r)]
    finally:
        wb.close()
    missing = [h for h in ("subcategoryId", "name", "orderIndex") if h not in headers]
    if missing:
        raise SystemExit(f"[-] {path.name}: missing column(s) {missing}")
    rows = [r for r in rows if str(r.get("name") or "").strip()]

    def order(r):
        try:
            return float(r.get("orderIndex"))
        except (TypeError, ValueError):
            return float("inf")

    rows.sort(key=order)
    if need_ids:
        empty = [r["name"] for r in rows if not str(r.get("subcategoryId") or "").strip()]
        if empty:
            raise SystemExit(f"[-] {path.name}: {len(empty)} row(s) without subcategoryId — this is the Phase-1 "
                             "file; upload it and use the platform export")
        ids = [str(r["subcategoryId"]).strip() for r in rows]
        dups = sorted({i for i in ids if ids.count(i) > 1})
        if dups:
            raise SystemExit(f"[-] {path.name}: duplicate subcategoryId {dups}")
    return rows


# --------------------------------------------------------------------------- fuzzy matching
def tokens(text: str) -> list[str]:
    text = ARABIC.sub(" ", TG_PREFIX.sub("", text)).casefold().replace("&", " and ")
    out: list[str] = []
    for t in re.split(r"[^a-z0-9]+", text):
        out += ABBREV[t].split() if t in ABBREV else [t]
    return [t for t in out if t and t not in STOP]


def _opposed(a: str, b: str) -> bool:
    """After the shared start, does one word continue with a meaning-changing prefix the other lacks?"""
    n = 0
    while n < min(len(a), len(b)) and a[n] == b[n]:
        n += 1
    for word, other in ((a, b), (b, a)):
        for pre in PREFIXES:
            # prefix straddling the shared start (hyp|er vs hyp|o) or starting right after it (hyper|para…)
            for k in range(max(0, n - len(pre) + 1), n + 1):
                if word.startswith(pre, k) and k + len(pre) > n and not other.startswith(pre, k) \
                        and (k == 0 or word[:k] == other[:k]):
                    return True
    return False


def _tok(a: str, b: str) -> float:
    if a == b:
        return 1.0
    if a.isdigit() or b.isdigit() or _opposed(a, b):
        return 0.0
    if min(len(a), len(b)) >= 4 and (a.startswith(b) or b.startswith(a)):
        return 0.9
    return SequenceMatcher(None, a, b).ratio()


def similarity(a: list[str], b: list[str]) -> float:
    """Typo-tolerant token dice (0.7) blended with a character ratio (0.3)."""
    if not a or not b:
        return 0.0

    def cover(x, y):
        total = 0.0
        for t in x:
            best = max(_tok(t, u) for u in y)
            total += best if best >= 0.75 else 0.0
        return total

    dice = (cover(a, b) + cover(b, a)) / (len(a) + len(b))
    seq = SequenceMatcher(None, " ".join(a), " ".join(b)).ratio()
    return round(0.7 * dice + 0.3 * seq, 3)


def _id_prefix(ids: list[str]) -> str:
    """Common category prefix of the ids ('DamiettaFa_X_2_'), cut at an underscore."""
    if len(ids) < 2:
        return ""
    pre = ids[0]
    for i in ids[1:]:
        while not i.startswith(pre):
            pre = pre[:-1]
    return pre[: pre.rfind("_") + 1] if "_" in pre else ""


def row_score(row: dict[str, Any], prefix: str, stem: str) -> float:
    sid = str(row["subcategoryId"])
    if stem.casefold() == sid.casefold():
        return 1.0
    ftoks = tokens(stem)
    best = similarity(tokens(str(row["name"])), ftoks)
    suffix = sid[len(prefix):] if prefix and sid.startswith(prefix) else sid
    return max(best, similarity(tokens(suffix), ftoks))


# --------------------------------------------------------------------------- sources
def parse_source_specs(specs: list[str], root: Path) -> list[tuple[Path, str | None]]:
    out = []
    for spec in specs:
        kind = None
        m = re.match(r"^([a-z_]+)=(.+)$", spec)
        if m and m.group(1) in KINDS:
            kind, spec = m.group(1), m.group(2)
        path = _abs(spec, root).resolve()
        if not path.exists():
            raise SystemExit(f"[-] source folder not found: {path}")
        out.append((path, kind))
    return out


def classify(path: Path, base: Path, forced: str | None) -> str:
    if forced:
        return forced
    if path.suffix.lower() in SLIDE_SUFFIXES:
        return "powerpoint"
    parts = [base.name] + list(path.relative_to(base).parts[:-1]) if path != base else [base.name]
    for part in reversed(parts):  # the folder closest to the file decides
        for kind, rx in DIR_KINDS:
            if rx.search(part):
                return kind
    return "other"


def scan_sources(specs: list[tuple[Path, str | None]]) -> list[dict[str, Any]]:
    files: list[dict[str, Any]] = []
    seen: set[Path] = set()
    for base, forced in specs:
        paths = [base] if base.is_file() else sorted(p for p in base.rglob("*") if p.is_file())
        for p in paths:
            if p.suffix.lower() not in LECTURE_SUFFIXES or p.name.startswith(("~$", ".")) or p.resolve() in seen:
                continue
            seen.add(p.resolve())
            files.append({"path": p.resolve(), "kind": classify(p, base if base.is_dir() else p.parent, forced),
                          "stem": p.stem, "key": " ".join(tokens(p.stem)) or p.stem.casefold()})
    return files


def _file_ref(ref: str, root: Path) -> dict[str, Any]:
    """'Book/x.pdf#31-32' → {"path": abs, "pages": "31-32"}."""
    path, _, pages = ref.partition("#")
    p = _abs(path, root).resolve()
    if not p.exists():
        raise SystemExit(f"[-] --map: file not found: {p}")
    return {"path": str(p), "pages": pages or None}


def load_map(path: str | None, root: Path, rows: list[dict[str, Any]]) -> dict[str, Any]:
    if not path:
        return {}
    data = json.loads(_abs(path, root).read_text(encoding="utf-8"))
    by_name = {str(r["name"]).strip().casefold(): str(r["subcategoryId"]) for r in rows}
    ids = {str(r["subcategoryId"]) for r in rows}
    out: dict[str, Any] = {}
    for key, spec in data.items():
        sid = key if key in ids else by_name.get(key.strip().casefold())
        if not sid:
            raise SystemExit(f"[-] --map: {key!r} is neither a subcategoryId nor a lecture name of the export")
        if spec is None:
            out[sid] = None
            continue
        if isinstance(spec, str):
            spec = {"files": [spec]}
        elif isinstance(spec, list):
            spec = {"files": spec}
        refs = [_file_ref(f, root) for f in spec.get("files", [])]
        if not refs:
            raise SystemExit(f"[-] --map: {key!r} has no files")
        kind = spec.get("kind") or classify(Path(refs[0]["path"]), Path(refs[0]["path"]).parent, None)
        out[sid] = {"files": refs, "kind": kind, "force": bool(spec.get("force"))}
    return out


def match(rows: list[dict[str, Any]], files: list[dict[str, Any]], priority: list[str], threshold: float,
          pins: dict[str, Any] | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """One-to-one greedy assignment per kind; each row takes the first kind in `priority` it has.

    A `--map` pin is a score-1.0 candidate of its kind (it replaces the row's fuzzy pick of that kind and
    still obeys the priority); `"force": true` or `null` (= missing) overrides everything."""
    pins = pins or {}
    pinned_paths = {f["path"] for pin in pins.values() if pin for f in pin["files"]}
    pool = [i for i, f in enumerate(files) if str(f["path"]) not in pinned_paths]
    prefix = _id_prefix([str(r["subcategoryId"]) for r in rows])
    scores: dict[tuple[str, int], float] = {}
    for r in rows:
        for i in pool:
            scores[(str(r["subcategoryId"]), i)] = row_score(r, prefix, files[i]["stem"])

    def blocked(sid: str, kind: str) -> bool:
        pin = pins.get(sid, False)
        return pin is None or (bool(pin) and (pin.get("force") or pin["kind"] == kind))

    assigned: dict[str, dict[str, int]] = {}      # kind → {sid: file index}
    for kind in {files[i]["kind"] for i in pool}:
        pairs = sorted(((s, sid, i) for (sid, i), s in scores.items()
                        if files[i]["kind"] == kind and s >= threshold and not blocked(sid, kind)),
                       key=lambda t: (-t[0], t[1], files[t[2]]["path"].suffix.lower() != ".pdf",
                                      str(files[t[2]]["path"])))
        taken_rows: dict[str, int] = {}
        taken_files: set[str] = set()
        for s, sid, i in pairs:
            if sid in taken_rows or files[i]["key"] in taken_files:
                continue
            taken_rows[sid] = i
            taken_files.add(files[i]["key"])
        assigned[kind] = taken_rows
    records = []
    for r in rows:
        sid = str(r["subcategoryId"])
        pin = pins.get(sid, False)
        rec = {"subcategoryId": sid, "name": r["name"], "tag": r.get("tag"), "orderIndex": r.get("orderIndex"),
               "status": "missing", "kind": None, "score": None, "files": [], "pinned": pin is not False,
               "alternatives": [], "ambiguous": False}
        ranked = sorted(((scores[(sid, i)], i) for i in pool if scores[(sid, i)] >= threshold), reverse=True)
        for kind in ([pin["kind"]] if pin and pin.get("force") else [] if pin is None else priority):
            if pin and pin["kind"] == kind:
                rec.update(status="matched", kind=kind, score=1.0, files=pin["files"])
                break
            i = assigned.get(kind, {}).get(sid)
            if i is None:
                continue
            rec.update(status="matched", kind=kind, score=scores[(sid, i)], pinned=False,
                       files=[{"path": str(files[i]["path"]), "pages": None}])
            # a rival is another file of the same kind that no other row took (a pdf/pptx pair of one deck
            # is the same lecture, not a rival)
            others = {files[j]["key"] for m in assigned.values() for s2, j in m.items() if s2 != sid}
            same = [s for s, j in ranked if files[j]["kind"] == kind and files[j]["key"] not in others
                    and files[j]["key"] != files[i]["key"]]
            rec["ambiguous"] = bool(same and rec["score"] - same[0] < 0.05 and rec["score"] < 1.0)
            break
        chosen = {f["path"] for f in rec["files"]}
        rec["alternatives"] = [{"path": str(files[i]["path"]), "kind": files[i]["kind"], "score": s}
                               for s, i in ranked if str(files[i]["path"]) not in chosen][:4]
        records.append(rec)
    # unmatched = no row wanted it in any kind and no pin names it (a --map candidate); lower-priority
    # alternatives and the pdf/pptx twin of a chosen deck are not reported
    wanted = {files[i]["key"] for m in assigned.values() for i in m.values()}
    unused = [files[i] for i in pool if files[i]["key"] not in wanted]
    return records, unused


# --------------------------------------------------------------------------- PDF helpers
def pdf_stats(path: Path) -> tuple[int | None, int]:
    """(pages, non-whitespace text-layer chars); (None, 0) when unreadable."""
    try:
        import fitz
        with fitz.open(path) as doc:
            chars = sum(len(re.sub(r"\s+", "", page.get_text())) for page in doc)
            return doc.page_count, chars
    except Exception:
        return None, 0


def convert_to_pdf(src: Path, workdir: Path) -> Path:
    if src.suffix.lower() == ".pdf":
        return src
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(f"{src.name}: LibreOffice (soffice) is needed to convert {src.suffix}")
    outdir = workdir / f"conv_{len(list(workdir.iterdir()))}"
    outdir.mkdir()
    profile = (workdir / "lo_profile").resolve()
    proc = subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless",
                           "--convert-to", "pdf", "--outdir", str(outdir), str(src)],
                          capture_output=True, text=True, timeout=600)
    out = outdir / f"{src.stem}.pdf"
    if proc.returncode or not out.exists():
        raise RuntimeError(f"LibreOffice could not convert {src.name}: {(proc.stderr or proc.stdout).strip()[:300]}")
    return out


def _pages(spec: str | None, total: int) -> list[int]:
    if not spec:
        return list(range(total))
    out: list[int] = []
    for part in str(spec).split(","):
        a, _, b = part.strip().partition("-")
        start, end = int(a), int(b or a)
        if start < 1 or end < start or end > total:
            raise RuntimeError(f"page range {part} outside 1-{total}")
        out.extend(range(start - 1, end))
    return out


def assemble(refs: list[dict[str, Any]], workdir: Path) -> Path:
    """Convert + slice + merge the refs into one PDF; a single whole PDF is returned as is."""
    pdfs = [(convert_to_pdf(Path(r["path"]), workdir), r.get("pages")) for r in refs]
    if len(pdfs) == 1 and not pdfs[0][1]:
        return pdfs[0][0]
    import fitz
    out = workdir / f"assembled_{len(list(workdir.iterdir()))}.pdf"
    with fitz.open() as dest:
        for path, pages in pdfs:
            with fitz.open(path) as doc:
                for p in _pages(pages, doc.page_count):
                    dest.insert_pdf(doc, from_page=p, to_page=p)
        dest.save(out)
    return out


def ocr_pdf(src: Path, workdir: Path) -> Path | None:
    exe = shutil.which("ocrmypdf")
    if not exe:
        return None
    out = workdir / f"ocr_{len(list(workdir.iterdir()))}.pdf"
    proc = subprocess.run([exe, "--skip-text", "-q", str(src), str(out)], capture_output=True, text=True)
    return out if proc.returncode == 0 and out.exists() else None


# --------------------------------------------------------------------------- reports
def _source_label(rec: dict[str, Any], root: Path) -> str:
    return " + ".join(_rel(Path(f["path"]), root) + (f"#{f['pages']}" if f.get("pages") else "")
                      for f in rec["files"]) or "—"


def write_match_report(meta: Path, manifest: dict[str, Any], root: Path, name: str = "lectures_match.md") -> Path:
    recs = manifest["lectures"]
    lines = [f"# {manifest['module']} — lecture sources", "",
             f"- Export: `{Path(manifest['export']).name}` · rows **{len(recs)}**",
             f"- Matched **{sum(r['status'] != 'missing' for r in recs)}** · missing "
             f"**{sum(r['status'] == 'missing' for r in recs)}** · ambiguous **{sum(r['ambiguous'] for r in recs)}**",
             f"- Priority: {' → '.join(manifest['priority'])} → missing · threshold {manifest['threshold']}", "",
             "| Order | Group | Lecture | Subcategory ID | Status | Kind | Score | Source | Pages | Text chars |",
             "|---:|---|---|---|---|---|---:|---|---:|---:|"]
    for r in recs:
        flag = " ⚠" if r["ambiguous"] else ""
        lines.append(f"| {r['orderIndex']} | {r.get('tag') or ''} | {r['name']} | `{r['subcategoryId']}` | "
                     f"{r['status']}{flag} | {r['kind'] or '—'} | {r['score'] if r['score'] is not None else '—'} | "
                     f"{_source_label(r, root)} | {r.get('pages') or '—'} | {r.get('textChars', '—')} |")
    if manifest.get("unusedSources"):
        lines += ["", f"## Unmatched source files ({len(manifest['unusedSources'])})", ""]
        lines += [f"- {u['kind']}: `{_rel(Path(u['path']), root)}`" for u in manifest["unusedSources"]]
    out = meta / "reports" / name
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out


# --------------------------------------------------------------------------- plan (Phase 1)
PART = re.compile(r"[\s_-]*(?:\(?\b(?:part|pt)\s*[-_.]?\s*(\d+|i{1,3}|iv|v)\b\)?|\b(i{1,3}|iv)\s*$)\s*$", re.I)


def _clean_name(text: str) -> str:
    text = ARABIC.sub(" ", TG_PREFIX.sub("", text))
    text = re.sub(r"^\s*(?:lecture|lec)?\s*\d+\s*[-_.)]+\s*", "", text, flags=re.I)
    text = re.sub(r"[_]+", " ", text)
    return re.sub(r"\s+", " ", text).strip(" -.")


def schedule_rows(path: Path, default_tag: str | None) -> list[tuple[str, str | None]]:
    """(name, tag) from .xlsx/.csv (name/lecture/title/topic + tag/group/subject columns) or .txt/.md
    (`# Group` / `Group:` headings, one lecture per line, list markers and numbering stripped)."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".csv"):
        if suffix == ".csv":
            import csv
            with open(path, newline="", encoding="utf-8-sig") as fh:
                table = [list(r) for r in csv.reader(fh)]
        else:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            table = [list(r) for r in wb.active.iter_rows(values_only=True)]
            wb.close()
        table = [r for r in table if any(v not in (None, "") for v in r)]
        head = [str(h or "").strip().casefold() for h in table[0]] if table else []
        pick = lambda names: next((i for i, h in enumerate(head) if h in names), None)  # noqa: E731
        ni = pick({"name", "lecture", "title", "topic", "lecture name"})
        ti = pick({"tag", "group", "subject", "department", "section"})
        body = table[1:] if ni is not None else table
        ni = ni or 0
        out = []
        for r in body:
            name = _clean_name(str(r[ni] or "")) if ni < len(r) else ""
            if name:
                tag = str(r[ti]).strip() if ti is not None and ti < len(r) and r[ti] else default_tag
                out.append((name, tag))
        return out
    tag, out = default_tag, []
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        m = re.match(r"^#+\s*(.+)$", line) or re.match(r"^([^-*•\d].*?):\s*$", line)
        if m:
            tag = _clean_name(m.group(1)) or tag
            continue
        name = _clean_name(re.sub(r"^(?:[-*•]|\d+[.)-])\s*", "", line))
        if name:
            out.append((name, tag))
    return out


def dir_rows(folder: Path, default_tag: str | None, merge_parts: bool) -> tuple[list[tuple[str, str | None]], list[list[str]]]:
    files = sorted((p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in LECTURE_SUFFIXES
                    and not p.name.startswith(("~$", "."))),
                   key=lambda p: [int(t) if t.isdigit() else t.casefold() for t in re.split(r"(\d+)", str(p.relative_to(folder)))])
    out: list[tuple[str, str | None]] = []
    parts: dict[tuple[str, str | None], list[str]] = {}
    for p in files:
        rel = p.relative_to(folder)
        tag = _clean_name(rel.parts[0]) if len(rel.parts) > 1 else default_tag
        name = _clean_name(p.stem)
        base = PART.sub("", name).strip() if PART.search(name) else None
        if base:
            parts.setdefault((base.casefold(), tag), []).append(name)
            if merge_parts:
                name = base
        if name and (name, tag) not in out and not any(n.casefold() == name.casefold() and t == tag for n, t in out):
            out.append((name, tag))
    groups = [names for names in parts.values() if len(names) > 1]
    return out, groups


def write_plan(out: Path, rows: list[tuple[str, str | None]], category: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for n, (name, tag) in enumerate(rows, 1):
        # Phase 1: categoryId and subcategoryId stay empty — the platform assigns them on upload.
        ws.append([None, category, None, name, tag, None, None, None, None, None, None, 0, n * 10])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F6B78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHIJKLM", (18, 20, 18, 78, 22, 12, 12, 24, 14, 14, 24, 16, 12)):
        ws.column_dimensions[col].width = width
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def cmd_plan(args) -> int:
    root = module_root(args.module)
    default_tag = args.tag
    groups: list[list[str]] = []
    if args.schedule:
        rows = schedule_rows(_abs(args.schedule, root), default_tag)
    else:
        rows, groups = dir_rows(_abs(args.from_dir, root), default_tag, args.merge_parts)
    if not rows:
        print("[-] no lecture found in the schedule/folder")
        return 1
    out = Path(args.out).resolve() if args.out else root / f"subcategories_{root.name}.xlsx"
    if out.exists() and not args.force:
        print(f"[-] {out} exists — pass --force to replace it")
        return 1
    category = args.category_name or root.name
    write_plan(out, rows, category)
    print(f"[+] {out}: {len(rows)} lecture(s), categoryName={category!r}, categoryId/subcategoryId empty\n")
    tag = object()
    for n, (name, t) in enumerate(rows, 1):
        if t != tag:
            tag = t
            print(f"  ## {t or '(no group)'} ({sum(1 for _, x in rows if x == t)})")
        print(f"   {n * 10:>5}  {name}")
    if groups:
        print(f"\n[!] {len(groups)} multi-part group(s) — decide merge vs separate with the user"
              f"{' (merged: --merge-parts)' if args.merge_parts else ' (re-run with --merge-parts to merge)'}:")
        for g in groups:
            print(f"    - {' | '.join(g)}")
    print(f"\n[STOP] Phase 1 done. Present this breakdown, then upload {out.name} to the platform and WAIT for "
          f"the export subcategories_{root.name}_<Date>.xlsx before `lectures {args.module} match`.")
    return 0


# --------------------------------------------------------------------------- match / apply / check
def _export_path(args, root: Path, manifest: dict[str, Any] | None = None) -> Path:
    if getattr(args, "export", None):
        return _abs(args.export, root).resolve()
    if manifest and manifest.get("export"):
        return Path(manifest["export"])
    return find_export(root)


def cmd_match(args) -> int:
    root = module_root(args.module)
    meta = meta_dir(root, args.out_dir)
    export = _export_path(args, root)
    rows = read_export(export)
    priority = [k.strip() for k in args.priority.split(",") if k.strip()]
    bad = [k for k in priority if k not in KINDS]
    if bad:
        raise SystemExit(f"[-] --priority: unknown kind(s) {bad}; kinds: {KINDS}")
    specs = parse_source_specs(args.sources, root)
    files = scan_sources(specs)
    pins = load_map(args.map, root, rows)
    records, unused = match(rows, files, priority, args.threshold, pins)
    # keep what a previous `apply` recorded for rows whose sources did not change, so a re-match does not
    # turn a non-byte-reproducible conversion/OCR output into an overwrite conflict
    prev_path = meta / MANIFEST
    prev = json.loads(prev_path.read_text(encoding="utf-8")) if prev_path.exists() else {}
    prev_by = {r["subcategoryId"]: r for r in prev.get("lectures", [])}
    for rec in records:
        old = prev_by.get(rec["subcategoryId"])
        if old and old.get("applied") and old.get("files") == rec["files"]:
            for k in ("applied", "output", "pages", "textChars"):
                if k in old:
                    rec[k] = old[k]
    manifest = {
        "generatedAt": _now(), "module": root.name, "moduleRoot": str(root), "export": str(export),
        "sources": [{"path": str(p), "kind": k or "auto"} for p, k in specs], "priority": priority,
        "threshold": args.threshold, "map": str(_abs(args.map, root).resolve()) if args.map else None,
        "lectureCount": len(records), "matchedCount": sum(r["status"] == "matched" for r in records),
        "missingCount": sum(r["status"] == "missing" for r in records),
        "lectures": records,
        "unusedSources": [{"path": str(u["path"]), "kind": u["kind"]} for u in unused],
    }
    meta.mkdir(parents=True, exist_ok=True)
    (meta / MANIFEST).write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    report = write_match_report(meta, manifest, root)
    by_kind: dict[str, int] = {}
    for r in records:
        by_kind[r["kind"] or "missing"] = by_kind.get(r["kind"] or "missing", 0) + 1
    for r in records:
        mark = "MISS" if r["status"] == "missing" else ("PIN " if r["pinned"] else "??  " if r["ambiguous"] else "OK  ")
        print(f"  {mark} {str(r['orderIndex']):>5} {r['name'][:55]:55} {r['kind'] or '':10} "
              f"{r['score'] if r['score'] is not None else '':>5}  {_source_label(r, root)[:70]}")
    print(f"\n[=] {len(records)} row(s): " + " · ".join(f"{k} {v}" for k, v in sorted(by_kind.items())) +
          f" · ambiguous {sum(r['ambiguous'] for r in records)} · unmatched files {len(unused)}")
    print(f"[=] manifest: {meta / MANIFEST}\n[=] report:   {report}")
    print("[=] dry run — nothing copied. Review the table (fix wrong rows with --map), then `lectures … apply`.")
    return 0


def _load_manifest(meta: Path) -> dict[str, Any]:
    path = meta / MANIFEST
    if not path.exists():
        raise SystemExit(f"[-] {path} not found — run `lectures <Module> match` first")
    return json.loads(path.read_text(encoding="utf-8"))


def cmd_apply(args) -> int:
    root = module_root(args.module)
    meta = meta_dir(root, args.out_dir)
    manifest = _load_manifest(meta)
    lectures = _abs(args.lectures_dir, root) if args.lectures_dir else root / "Lectures"
    lectures.mkdir(parents=True, exist_ok=True)
    counts = {"written": 0, "unchanged": 0, "conflict": 0, "error": 0, "missing": 0}
    with tempfile.TemporaryDirectory(prefix="mbset_lectures_") as tmp:
        work = Path(tmp)
        for rec in manifest["lectures"]:
            dest = lectures / f"{rec['subcategoryId']}.pdf"
            if rec["status"] == "missing" or not rec["files"]:
                counts["missing"] += 1
                continue
            try:
                src_sha = "+".join(sha256(Path(f["path"]))[:16] + (f"#{f['pages']}" if f.get("pages") else "")
                                   for f in rec["files"])
                prev = rec.get("applied") or {}
                if (dest.exists() and prev.get("sourceSha") == src_sha and prev.get("ocr", False) == args.ocr
                        and prev.get("outputSha") == sha256(dest) and not args.force):
                    counts["unchanged"] += 1
                    continue
                built = assemble(rec["files"], work)
                ocred = False
                if args.ocr:
                    _, chars = pdf_stats(built)
                    if chars < args.min_chars:
                        done = ocr_pdf(built, work)
                        if done is None:
                            print(f"  [!] {rec['subcategoryId']}: ocrmypdf unavailable/failed — kept without OCR")
                        else:
                            built, ocred = done, True
                if dest.exists():
                    same = sha256(dest) == sha256(built)
                    if same:
                        counts["unchanged"] += 1
                    elif not args.force:
                        counts["conflict"] += 1
                        print(f"  [-] {dest.name}: exists with different content — not overwritten (use --force)")
                        continue
                if not dest.exists() or sha256(dest) != sha256(built):
                    part = dest.with_name(dest.name + ".part")
                    shutil.copy2(built, part)
                    part.replace(dest)
                    counts["written"] += 1
                    print(f"  [+] {dest.name} ← {_source_label(rec, root)[:80]}{' (OCR)' if ocred else ''}")
                pages, chars = pdf_stats(dest)
                rec.update(status="ready", output=str(dest), pages=pages, textChars=chars,
                           applied={"at": _now(), "sourceSha": src_sha, "outputSha": sha256(dest),
                                    "ocr": args.ocr, "ocrApplied": ocred})
            except Exception as exc:  # noqa: BLE001 — one bad source must not stop the others
                counts["error"] += 1
                rec["error"] = f"{type(exc).__name__}: {exc}"
                print(f"  [-] {rec['subcategoryId']}: {rec['error']}")
    manifest["appliedAt"], manifest["lecturesDir"] = _now(), str(lectures)
    (meta / MANIFEST).write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    report = write_match_report(meta, manifest, root)
    missing = [r for r in manifest["lectures"] if r["status"] == "missing"]
    print(f"\n[=] written {counts['written']} · unchanged {counts['unchanged']} · conflict {counts['conflict']} · "
          f"error {counts['error']} · missing {counts['missing']}")
    for r in missing:
        print(f"    missing: {r['orderIndex']} {r['name']} ({r['subcategoryId']})")
    print(f"[=] report: {report}")
    return 1 if counts["conflict"] or counts["error"] else 0


def cmd_check(args) -> int:
    root = module_root(args.module)
    meta = meta_dir(root, args.out_dir)
    manifest = json.loads((meta / MANIFEST).read_text(encoding="utf-8")) if (meta / MANIFEST).exists() else None
    export = _export_path(args, root, manifest)
    rows = read_export(export)
    lectures = _abs(args.lectures_dir, root) if args.lectures_dir else root / "Lectures"
    ids = {str(r["subcategoryId"]) for r in rows}
    present = sorted(p for p in lectures.iterdir() if p.is_file()) if lectures.is_dir() else []
    stray = [p.name for p in present if p.suffix.lower() != ".pdf" or p.stem not in ids]
    missing, unreadable, low = [], [], []
    lines = [f"# {root.name} — lectures check", "", f"- Export: `{export.name}` · rows **{len(rows)}**",
             f"- Folder: `{_rel(lectures, root)}`", "",
             "| Order | Lecture | Subcategory ID | PDF | Pages | Text chars |", "|---:|---|---|---|---:|---:|"]
    for r in rows:
        sid = str(r["subcategoryId"])
        path = lectures / f"{sid}.pdf"
        if not path.exists():
            missing.append(r)
            lines.append(f"| {r['orderIndex']} | {r['name']} | `{sid}` | MISSING | — | — |")
            continue
        pages, chars = pdf_stats(path)
        if not pages:
            unreadable.append(sid)
        elif chars < args.min_chars:
            low.append(sid)
        lines.append(f"| {r['orderIndex']} | {r['name']} | `{sid}` | ok | {pages if pages else 'unreadable'} | {chars} |")
        if args.verbose:
            print(f"  {str(r['orderIndex']):>5} {sid[:70]:70} {pages or '?':>4}p {chars:>7} chars")
    if stray:
        lines += ["", f"## Stray files ({len(stray)})", ""] + [f"- `{s}`" for s in stray]
    out = meta / "reports" / "lectures_check.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    ok = len(rows) - len(missing)
    print(f"[{'+' if not (missing or stray or unreadable) else '-'}] {ok}/{len(rows)} row(s) have "
          f"{_rel(lectures, root)}/<subcategoryId>.pdf")
    for r in missing:
        print(f"    missing: {r['orderIndex']} {r['name']} ({r['subcategoryId']})")
    for s in stray:
        print(f"    stray:   {s}")
    for s in unreadable:
        print(f"    unreadable: {s}.pdf")
    if low:
        print(f"[!] {len(low)} PDF(s) with < {args.min_chars} text-layer chars (scanned — `apply --ocr` adds one): "
              + ", ".join(low[:8]) + (" …" if len(low) > 8 else ""))
    print(f"[=] report: {out}")
    return 1 if (missing or stray or unreadable) else 0


# --------------------------------------------------------------------------- argparse
def run(args) -> int:
    return {"plan": cmd_plan, "match": cmd_match, "apply": cmd_apply, "check": cmd_check}[args.lectures_cmd](args)


def register(sub) -> None:
    p = sub.add_parser("lectures", help="Phase 1/2 lectures: plan subcategories, match sources, apply, check")
    p.add_argument("module", help="module folder, e.g. 'My University/Endocrinology'")
    acts = p.add_subparsers(dest="lectures_cmd", required=True)

    def common(q, export=True):
        q.add_argument("--out-dir", help="manifest/report folder instead of <Module>/.mbset")
        if export:
            q.add_argument("--export", help="platform export (default: newest subcategories_<Module>_*.xlsx)")

    q = acts.add_parser("plan", help="Phase 1: write subcategories_<Module>.xlsx with empty ids, then STOP")
    src = q.add_mutually_exclusive_group(required=True)
    src.add_argument("--schedule", help=".xlsx/.csv (name[, tag] columns) or .txt/.md (# Group headings)")
    src.add_argument("--from-dir", help="one row per lecture file; first sub-folder = tag")
    q.add_argument("--category-name", help="default: the module folder name")
    q.add_argument("--tag", help="tag for rows without a group")
    q.add_argument("--merge-parts", action="store_true", help="collapse 'X Part 1/2' files into one row")
    q.add_argument("--out", help="default: <Module>/subcategories_<Module>.xlsx")
    q.add_argument("--force", action="store_true")
    q = acts.add_parser("match", help="Phase 2 dry run: official row → lecture source; writes manifest + report")
    common(q)
    q.add_argument("--sources", nargs="+", required=True, metavar="[KIND=]DIR",
                   help=f"source folders/files; KIND one of {KINDS} (default: from extension/folder name)")
    q.add_argument("--priority", default=",".join(DEFAULT_PRIORITY),
                   help="kinds in preference order; unlisted kinds are never used")
    q.add_argument("--threshold", type=float, default=0.6, help="minimum fuzzy score (0-1)")
    q.add_argument("--map", help="JSON of pinned sources: book page ranges, merges, null = missing")
    q = acts.add_parser("apply", help="copy/convert matched sources to Lectures/<subcategoryId>.pdf")
    common(q, export=False)
    q.add_argument("--lectures-dir", help="default: <Module>/Lectures")
    q.add_argument("--ocr", action="store_true", help="ocrmypdf outputs without a text layer")
    q.add_argument("--min-chars", type=int, default=200, help="text-layer chars below which a PDF counts as scanned")
    q.add_argument("--force", action="store_true", help="overwrite an existing different Lectures/<id>.pdf")
    q = acts.add_parser("check", help="every row has Lectures/<id>.pdf, no strays; pages + text layer")
    common(q)
    q.add_argument("--lectures-dir", help="default: <Module>/Lectures")
    q.add_argument("--min-chars", type=int, default=200)
    q.add_argument("--verbose", action="store_true", help="print one line per PDF")
    p.set_defaults(fn=run)
